from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "GUIA_RAPIDO_DESKTOP_IMPACT.docx"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margin(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_step(document: Document, number: int, title: str, detail: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.left_indent = Inches(0.08)
    run = paragraph.add_run(f"{number}. {title}  ")
    run.bold = True
    run.font.color.rgb = RGBColor(35, 35, 35)
    paragraph.add_run(detail)


document = Document()
section = document.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.62)
section.bottom_margin = Inches(0.58)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = document.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(22)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
styles["Title"].paragraph_format.space_after = Pt(8)
for style_name, size in (("Heading 1", 14), ("Heading 2", 11.5)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(5)

title = document.add_paragraph(style="Title")
title.add_run("Guia de abertura e atualizacao do Estudo Desktop")

intro = document.add_paragraph()
intro.add_run("Objetivo. ").bold = True
intro.add_run(
    "Este guia mostra como abrir o dashboard no computador do gerente e como carregar novas bases sem alterar os arquivos originais."
)

document.add_heading("Como abrir", level=1)
add_step(document, 1, "Acesse a pasta do projeto.", r"C:\OneDrive - Claro SA\Jonathan\Desktop\ESTUDO_DESKTOP_OUTPUT\desktop-impact-dashboard")
add_step(document, 2, "Execute o aplicativo.", "Clique duas vezes em DesktopImpact.exe.")
add_step(document, 3, "Aguarde a abertura.", "O painel sera exibido em uma janela propria; nenhuma instalacao de Python e necessaria.")
add_step(document, 4, "Encerre corretamente.", "Feche a janela do dashboard para finalizar o servidor local.")

document.add_heading("Como atualizar as bases", level=1)
paragraph = document.add_paragraph("Mantenha as fontes nos caminhos abaixo, com estes nomes exatos:")

for filename in (
    r"ESTUDO_DESKTOP_OUTPUT\input_ESTUDO_DESKTOP_V3xlsx.xlsx",
    r"ESTUDO_DESKTOP_OUTPUT\input_BASE_CIDADES_APOIO.xlsx",
    "Senha_Ago.csv",
    "Fluxo.csv",
):
    item = document.add_paragraph(style="List Bullet")
    item.paragraph_format.space_after = Pt(1)
    item.add_run(filename)

add_step(document, 1, "Inicie a carga.", "Execute Atualizar_Dashboard.bat na pasta do projeto.")
add_step(document, 2, "Aguarde o processamento.", "A tela de atualizacao permanecera aberta ate o painel ser reconstruido.")
add_step(document, 3, "Confira o painel.", "A janela sera aberta automaticamente com os novos arquivos derivados.")

note = document.add_paragraph()
note.add_run("Importante. ").bold = True
note.add_run("As fontes sao apenas lidas. Copias de trabalho ficam em source_working e as saidas do dashboard ficam em data.")

document.add_heading("Arquivos principais", level=1)
table = document.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.columns[0].width = Inches(2.35)
table.columns[1].width = Inches(4.55)
headers = table.rows[0].cells
headers[0].text = "Arquivo ou pasta"
headers[1].text = "Funcao"
set_repeat_table_header(table.rows[0])
for cell in headers:
    shade(cell, "303030")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margin(cell)
    for run in cell.paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

rows = (
    ("DesktopImpact.exe", "Abre o dashboard sem exigir Python."),
    ("Atualizar_Dashboard.bat", "Le as fontes, recria os dados derivados e abre o painel."),
    ("dashboard_config.json", "Define a pasta das fontes, os nomes esperados e a porta inicial."),
    ("tools\\prepare_data_v3.py", "Contem as regras de leitura, validacao e calculo."),
    ("data", "Armazena JSON e CSV derivados usados pelo dashboard."),
)
for index, (name, purpose) in enumerate(rows):
    cells = table.add_row().cells
    cells[0].text = name
    cells[1].text = purpose
    if index % 2:
        for cell in cells:
            shade(cell, "F3F4F6")
    for cell in cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margin(cell)

document.add_heading("Cuidados rapidos", level=1)
for text in (
    "Nao mova somente o executavel; mantenha todos os arquivos e pastas do projeto juntos.",
    "O mapa-base precisa de internet para carregar os blocos cartograficos.",
    "Se o Windows bloquear o executavel, siga a politica de seguranca da Claro ou acione o suporte de TI.",
):
    item = document.add_paragraph(style="List Bullet")
    item.paragraph_format.space_after = Pt(2)
    item.add_run(text)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("Estudo Desktop | Canal AA e Vivo | Guia rapido")
run.font.name = "Aptos"
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(100, 100, 100)

document.save(OUTPUT)
print(OUTPUT)
